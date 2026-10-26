import os
import sys
from datetime import datetime
from tkinter import Frame, Label, messagebox, StringVar
import tkinter.ttk as ttk

from .db import get_db_connection
from .utils import fmt_date, fmt_qty_clean, get_desktop_path


class BalancesTabMixin:

    def _build_balances_tab(self):
        frame = self.tab_balances

        info_frame = Frame(frame, bg=self.bg_panel)
        info_frame.pack(fill="x", padx=20, pady=18)
        info_frame.grid_columnconfigure(0, weight=1)

        Label(info_frame, text="Остатки товаров на складе",
              bg=self.bg_panel, fg=self.fg_text,
              font=("Segoe UI", 16, "bold")).grid(
            row=0, column=0, sticky="w", pady=(0, 2))
        Label(info_frame, text="Только товары с проведёнными операциями",
              bg=self.bg_panel, fg=self.fg_light,
              font=("Segoe UI", 11)).grid(row=1, column=0, sticky="w")

        has_excel_icon = self.icons.get("excel") is not None
        excel_text = " Сохранить в Excel" if has_excel_icon else "📊 Сохранить в Excel"
        self.export_excel_btn = self.make_icon_button(
            info_frame, excel_text,
            "excel",
            self.export_balances_to_excel,
            self.color_green,
            fg=self.fg_dark,
            padx=18, pady=8)
        self.export_excel_btn.grid(row=0, column=1, rowspan=2,
                                   sticky="e", padx=(10, 0))

        # --- НОВОЕ: поиск по остаткам ----------------------------------
        search_frame = Frame(frame, bg=self.bg_panel)
        search_frame.pack(fill="x", padx=20, pady=(0, 10))

        Label(search_frame, text="🔍 Поиск:", bg=self.bg_panel, fg=self.fg_text,
              font=("Segoe UI", 11)).pack(side="left", padx=(0, 8))
        self.balance_search_var = StringVar()
        self.balance_search_entry = ttk.Entry(
            search_frame, textvariable=self.balance_search_var,
            width=40, style="Custom.TEntry")
        self.balance_search_entry.pack(side="left", fill="x", expand=True)
        self.bind_copy_paste(self.balance_search_entry)
        self.balance_search_entry.bind("<Return>",
                                       lambda e: self.update_balances_table())
        self.balance_search_entry.bind("<KeyRelease>",
                                       lambda e: self.update_balances_table())

        table_frame = Frame(frame, bg=self.bg_panel)
        table_frame.pack(fill="both", expand=True, padx=20, pady=(0, 10))

        columns = ("product", "gtin", "income", "outcome",
                   "balance", "last_in", "last_out")
        self.balances_tree = ttk.Treeview(table_frame, columns=columns, show="headings")

        self.balances_tree.heading("product", text="Товар", anchor="center")
        self.balances_tree.heading("gtin", text="GTIN", anchor="center")
        self.balances_tree.heading("income", text="Приход", anchor="center")
        self.balances_tree.heading("outcome", text="Расход", anchor="center")
        self.balances_tree.heading("balance", text="Остаток", anchor="center")
        self.balances_tree.heading("last_in", text="Посл. приход", anchor="center")
        self.balances_tree.heading("last_out", text="Посл. расход", anchor="center")

        self.balances_tree.column("product", width=300, anchor="w")
        self.balances_tree.column("gtin", width=150, anchor="center")
        self.balances_tree.column("income", width=80, anchor="center")
        self.balances_tree.column("outcome", width=80, anchor="center")
        self.balances_tree.column("balance", width=80, anchor="center")
        self.balances_tree.column("last_in", width=120, anchor="center")
        self.balances_tree.column("last_out", width=120, anchor="center")

        self.balances_tree.tag_configure('odd', background=self.bg_panel)
        self.balances_tree.tag_configure('even', background=self.bg_input)
        self.balances_tree.tag_configure('positive', foreground=self.color_yellow)
        self.balances_tree.tag_configure('negative', foreground=self.color_red)
        self.balances_tree.tag_configure('zero', foreground=self.fg_light)

        scroll_y = ttk.Scrollbar(table_frame, orient="vertical",
                                 command=self.balances_tree.yview)
        self.balances_tree.configure(yscrollcommand=scroll_y.set)

        self.balances_tree.pack(side="left", fill="both", expand=True)
        scroll_y.pack(side="right", fill="y")

        self.balances_tree.bind("<ButtonRelease-1>", self.show_product_operations)

    def update_balances_table(self, search=None):
        if search is None:
            search = (self.balance_search_var.get().strip().lower()
                      if hasattr(self, "balance_search_var") else "")

        for item in self.balances_tree.get_children():
            self.balances_tree.delete(item)

        conn = get_db_connection()
        cur = conn.cursor()
        query = """
            SELECT
                p.id,
                p.name,
                p.gtin,
                COALESCE(SUM(CASE WHEN m.mtype='in' THEN m.qty ELSE 0 END), 0) AS total_in,
                COALESCE(SUM(CASE WHEN m.mtype='out' THEN m.qty ELSE 0 END), 0) AS total_out,
                COALESCE(SUM(CASE WHEN m.mtype='in' THEN m.qty ELSE -m.qty END), 0) AS balance,
                MAX(CASE WHEN m.mtype='in' THEN m.date END) AS last_in,
                MAX(CASE WHEN m.mtype='out' THEN m.date END) AS last_out
            FROM products p
            LEFT JOIN movements m ON p.id = m.product_id
        """
        params = []
        if search:
            query += " WHERE LOWER(p.name) LIKE ? OR LOWER(COALESCE(p.gtin,'')) LIKE ?"
            params = [f"%{search}%", f"%{search}%"]
        query += """
            GROUP BY p.id, p.name, p.gtin
            HAVING total_in > 0 OR total_out > 0
            ORDER BY p.sort_order, p.name
        """
        cur.execute(query, params)
        rows = cur.fetchall()
        conn.close()

        for idx, row in enumerate(rows):
            tag = 'even' if idx % 2 == 0 else 'odd'
            balance = row["balance"]
            if balance > 0:
                balance_tag = 'positive'
            elif balance < 0:
                balance_tag = 'negative'
            else:
                balance_tag = 'zero'

            self.balances_tree.insert(
                "", "end",
                iid=str(row["id"]),
                values=(row["name"],
                       row["gtin"] or "-",
                       fmt_qty_clean(row["total_in"]),
                       fmt_qty_clean(row["total_out"]),
                       fmt_qty_clean(balance),
                       fmt_date(row["last_in"]),
                       fmt_date(row["last_out"])),
                tags=(tag, balance_tag))

    def show_product_operations(self, event):
        selection = self.balances_tree.selection()
        if not selection:
            return

        product_id = int(selection[0])

        # Сбрасываем поиск, чтобы пользователь видел полную историю товара
        if hasattr(self, "movement_search_var"):
            self.movement_search_var.set("")

        self.notebook.select(self.tab_movements)

        for item in self.movements_tree.get_children():
            self.movements_tree.delete(item)

        conn = get_db_connection()
        cur = conn.cursor()
        cur.execute("""
            SELECT m.id, m.date, m.qty, m.mtype, p.name, p.gtin
            FROM movements m
            JOIN products p ON m.product_id = p.id
            WHERE m.product_id = ?
            ORDER BY m.date DESC, m.id DESC
        """, (product_id,))
        rows = cur.fetchall()
        conn.close()

        for idx, row in enumerate(rows):
            type_text = "Приход" if row["mtype"] == "in" else "Расход"
            tag = 'even' if idx % 2 == 0 else 'odd'
            self.movements_tree.insert(
                "", "end",
                values=(row["name"], row["gtin"] or "", type_text,
                       fmt_qty_clean(row["qty"]), fmt_date(row["date"])),
                iid=str(row["id"]), tags=(tag,))

        self.op_status.config(
            text="Показаны операции по выбранному товару. "
                 "Нажмите «Отмена», чтобы сбросить фильтр.",
            fg=self.color_yellow)

    # export_balances_to_excel остаётся без изменений
    def export_balances_to_excel(self):
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        except ImportError:
            messagebox.showerror(
                "Ошибка",
                "Для экспорта в Excel требуется библиотека openpyxl.\n\n"
                "Установите её командой:\n    pip install openpyxl")
            return

        desktop = get_desktop_path()
        folder = os.path.join(desktop, "Учёт остатков")
        try:
            os.makedirs(folder, exist_ok=True)
        except Exception as e:
            messagebox.showerror(
                "Ошибка",
                f"Не удалось создать папку:\n{folder}\n\n{e}")
            return

        filepath = os.path.join(folder, "Остатки.xlsx")

        conn = get_db_connection()
        cur = conn.cursor()
        cur.execute("""
            SELECT
                p.id, p.name, p.gtin,
                COALESCE(SUM(CASE WHEN m.mtype='in' THEN m.qty ELSE 0 END), 0) AS total_in,
                COALESCE(SUM(CASE WHEN m.mtype='out' THEN m.qty ELSE 0 END), 0) AS total_out,
                COALESCE(SUM(CASE WHEN m.mtype='in' THEN m.qty ELSE -m.qty END), 0) AS balance,
                MAX(CASE WHEN m.mtype='in' THEN m.date END) AS last_in,
                MAX(CASE WHEN m.mtype='out' THEN m.date END) AS last_out
            FROM products p
            LEFT JOIN movements m ON p.id = m.product_id
            GROUP BY p.id, p.name, p.gtin
            HAVING total_in > 0 OR total_out > 0
            ORDER BY p.sort_order, p.name
        """)
        rows = cur.fetchall()
        conn.close()

        thin = Side(style="thin", color="CCCCCC")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        header_fill = PatternFill("solid", fgColor="249E80")
        alt_fill = PatternFill("solid", fgColor="F2F2F2")

        wb = Workbook()
        ws = wb.active
        ws.title = "Остатки"

        ws.merge_cells("A1:G1")
        ws["A1"] = "Отчёт по остаткам на складе"
        ws["A1"].font = Font(bold=True, size=14)
        ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 24

        ws.merge_cells("A2:G2")
        ws["A2"] = f"Дата формирования: {datetime.now().strftime('%d.%m.%Y %H:%M')}"
        ws["A2"].font = Font(italic=True, size=10, color="666666")
        ws["A2"].alignment = Alignment(horizontal="center")

        headers = ["Товар", "GTIN", "Приход", "Расход", "Остаток",
                   "Посл. приход", "Посл. расход"]
        header_row = 4
        for col, h in enumerate(headers, 1):
            c = ws.cell(row=header_row, column=col, value=h)
            c.font = Font(bold=True, color="FFFFFF", size=11)
            c.fill = header_fill
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border = border
        ws.row_dimensions[header_row].height = 22

        for i, row in enumerate(rows):
            r = header_row + 1 + i
            ws.cell(row=r, column=1, value=row["name"])
            ws.cell(row=r, column=2, value=row["gtin"] or "")
            ws.cell(row=r, column=3, value=float(row["total_in"]))
            ws.cell(row=r, column=4, value=float(row["total_out"]))
            ws.cell(row=r, column=5, value=float(row["balance"]))
            ws.cell(row=r, column=6, value=fmt_date(row["last_in"]) if row["last_in"] else "-")
            ws.cell(row=r, column=7, value=fmt_date(row["last_out"]) if row["last_out"] else "-")

            for col in range(1, 8):
                cell = ws.cell(row=r, column=col)
                cell.border = border
                if i % 2 == 1:
                    cell.fill = alt_fill
                if col in (3, 4, 5):
                    cell.number_format = "0.00"
                    cell.alignment = Alignment(horizontal="right")
                elif col == 1:
                    cell.alignment = Alignment(horizontal="left")
                else:
                    cell.alignment = Alignment(horizontal="center")

        max_len_a = max([len(str(row["name"])) for row in rows] + [20])
        ws.column_dimensions["A"].width = min(max_len_a + 5, 80)

        for col_letter, width in {"B": 20, "C": 12, "D": 12,
                                   "E": 12, "F": 15, "G": 15}.items():
            ws.column_dimensions[col_letter].width = width

        ws.freeze_panes = "A5"

        try:
            wb.save(filepath)
        except PermissionError:
            messagebox.showerror(
                "Файл занят",
                f"Не удалось перезаписать файл:\n{filepath}\n\n"
                "Возможно, он открыт в Excel. Закройте файл и повторите попытку.")
            return
        except Exception as e:
            messagebox.showerror("Ошибка", f"Не удалось сохранить файл:\n{e}")
            return

        if messagebox.askyesno(
                "Готово",
                f"Отчёт сохранён:\n{filepath}\n\nОткрыть файл?"):
            try:
                if os.name == "nt":
                    os.startfile(filepath)
                elif sys.platform == "darwin":
                    import subprocess
                    subprocess.Popen(["open", filepath])
                else:
                    import subprocess
                    subprocess.Popen(["xdg-open", filepath])
            except Exception as e:
                messagebox.showerror("Ошибка", f"Не удалось открыть файл:\n{e}")